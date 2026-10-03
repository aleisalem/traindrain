"""Document import: Markdown, DOCX, and PDF converted into a fresh draft module.

Mirrors `test_module_transfer.py`'s posture: the happy path per format is one
test each, and the rest is the input being treated as hostile — a remote
image that must never be fetched, raw HTML/script text that must never become
an executable node, a `.docx` zip bomb, a renamed spreadsheet, and an
encrypted PDF.
"""

import io
import uuid
import zipfile
from typing import Any

from docx import Document
from docx.shared import Inches
from httpx import AsyncClient
from PIL import Image
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import AuditLog, Module, ModuleAsset, ModulePage, Role, User
from app.security.passwords import hash_password
from app.security.sessions import COOKIE_NAME

KNOWN_PASSWORD = "a-perfectly-fine-passphrase"

_PDF_ONE_PAGE = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Resources<</Font<</F1 4 0 R>>>>/Contents 5 0 R>>endobj
4 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
5 0 obj<</Length 44>>
stream
BT /F1 24 Tf 10 100 Td (Hello world) Tj ET
endstream
endobj
trailer<</Size 6/Root 1 0 R>>
startxref
0
%%EOF"""


async def _make_user_with_role(db_session: AsyncSession, *, email: str, role_name: str) -> User:
    role = (await db_session.execute(select(Role).where(Role.name == role_name))).scalar_one()
    user = User(
        id=uuid.uuid4(),
        email=email,
        password_hash=hash_password(KNOWN_PASSWORD),
        first_name="Cora",
        last_name="Manager",
        roles=[role],
    )
    db_session.add(user)
    await db_session.commit()
    return user


async def _login(client: AsyncClient, *, email: str) -> None:
    response = await client.post(
        "/api/auth/login", json={"email": email, "password": KNOWN_PASSWORD}
    )
    client.cookies.set(COOKIE_NAME, response.cookies[COOKIE_NAME])


async def _login_with_role(
    client: AsyncClient, db_session: AsyncSession, *, email: str, role_name: str
) -> User:
    user = await _make_user_with_role(db_session, email=email, role_name=role_name)
    await _login(client, email=email)
    return user


async def _author(client: AsyncClient, db_session: AsyncSession, *, email: str) -> User:
    return await _login_with_role(client, db_session, email=email, role_name="Content Manager")


async def _import_document(
    client: AsyncClient,
    data: bytes,
    *,
    filename: str,
    language: str = "en",
    content_type: str = "application/octet-stream",
) -> Any:
    return await client.post(
        "/api/content/modules/document-import",
        data={"language": language},
        files={"file": (filename, data, content_type)},
    )


async def _module_pages(db_session: AsyncSession, module_id: str) -> list[ModulePage]:
    return list(
        (
            await db_session.execute(
                select(ModulePage)
                .where(ModulePage.module_id == uuid.UUID(module_id))
                .order_by(ModulePage.position)
            )
        ).scalars()
    )


def _report_codes(body: dict[str, Any]) -> set[str]:
    return {entry["code"] for entry in body["conversion_report"]}


def _png_bytes(*, color: str = "red", size: tuple[int, int] = (4, 4)) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", size, color=color).save(buffer, format="PNG")
    return buffer.getvalue()


def _build_docx(*, with_table: bool = False, with_image: bool = False) -> bytes:
    doc = Document()
    doc.add_heading("Imported Title", level=1)
    doc.add_heading("Getting Started", level=2)
    paragraph = doc.add_paragraph("Welcome to the training. ")
    bold_run = paragraph.add_run("Read carefully.")
    bold_run.bold = True
    doc.add_paragraph("First bullet", style="List Bullet")
    doc.add_paragraph("Second bullet", style="List Bullet")
    doc.add_paragraph("First step", style="List Number")
    doc.add_paragraph("Second step", style="List Number")
    if with_table:
        doc.add_table(rows=2, cols=2)
    if with_image:
        image_paragraph = doc.add_paragraph()
        image_stream = io.BytesIO(_png_bytes())
        image_paragraph.add_run().add_picture(image_stream, width=Inches(0.5))
    doc.add_heading("Next Section", level=2)
    doc.add_paragraph("Closing paragraph.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_zip(entries: dict[str, bytes | str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return buffer.getvalue()


_SPREADSHEET_CONTENT_TYPES = (
    '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Override PartName="/xl/workbook.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
    "</Types>"
)

_WORD_CONTENT_TYPES = (
    '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Override PartName="/word/document.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)


# --- Markdown ----------------------------------------------------------------


async def test_markdown_import_creates_module_with_pages(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="author1@example.com")
    markdown = (
        "# Awareness Basics\n\n"
        "Intro paragraph with **bold** and *italic* and `code`.\n\n"
        "## Spotting a Phish\n\n"
        "- Check the sender\n"
        "- Check the links\n\n"
        "1. Pause\n"
        "2. Verify\n\n"
        "> When in doubt, report it.\n\n"
        "```python\nprint('safe')\n```\n\n"
        "---\n\n"
        "## Reporting\n\n"
        "Use the button in your mail client.\n"
    )
    response = await _import_document(client, markdown.encode("utf-8"), filename="awareness.md")
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["module"]["title"] == "Awareness Basics"
    assert body["module"]["status"] == "draft"
    assert body["module"]["language"] == "en"

    pages = await _module_pages(db_session, body["module"]["id"])
    # The intro paragraph before the first `## ` heading becomes its own
    # leading page, titled after the module itself — content is never
    # dropped just for lacking a heading of its own.
    assert [page.title for page in pages] == ["Awareness Basics", "Spotting a Phish", "Reporting"]

    intro_paragraph = pages[0].body["content"][0]
    marks_seen = {
        mark["type"] for node in intro_paragraph["content"] for mark in node.get("marks", [])
    }
    assert marks_seen == {"bold", "italic", "code"}

    second_page_types = [node["type"] for node in pages[1].body["content"]]
    assert "bulletList" in second_page_types
    assert "orderedList" in second_page_types
    assert "blockquote" in second_page_types
    assert "codeBlock" in second_page_types
    assert "horizontalRule" in second_page_types


async def test_markdown_remote_image_dropped_and_reported(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="author2@example.com")
    markdown = (
        "# Illustrated Guide\n\n"
        "## Section One\n\n"
        "Look at this: ![diagram](https://evil.example.com/tracker.png)\n"
    )
    response = await _import_document(client, markdown.encode("utf-8"), filename="guide.md")
    assert response.status_code == 201, response.text
    body = response.json()
    assert "image_reference_dropped" in _report_codes(body)

    assets = (
        await db_session.execute(
            select(ModuleAsset).where(ModuleAsset.module_id == uuid.UUID(body["module"]["id"]))
        )
    ).scalars().all()
    assert assets == []

    pages = await _module_pages(db_session, body["module"]["id"])
    serialized = str(pages[0].body)
    assert "evil.example.com" not in serialized
    assert "image" not in {node["type"] for node in pages[0].body["content"]}


async def test_markdown_raw_html_is_never_stored_as_a_node(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="author3@example.com")
    markdown = "# Notes\n\n## Body\n\n<script>alert(1)</script>\n\nSome <b>raw</b> html here.\n"
    response = await _import_document(client, markdown.encode("utf-8"), filename="notes.md")
    assert response.status_code == 201, response.text
    body = response.json()

    pages = await _module_pages(db_session, body["module"]["id"])
    allowed_node_types = {
        "doc", "paragraph", "heading", "bulletList", "orderedList", "listItem",
        "blockquote", "codeBlock", "horizontalRule", "image", "text",
    }

    def _walk(node: dict[str, Any]) -> None:
        assert node["type"] in allowed_node_types
        for child in node.get("content") or []:
            _walk(child)

    for page in pages:
        _walk(page.body)

    # markdown-it never even tokenizes this as HTML with `html: False` — it
    # becomes ordinary text, present verbatim but inert, never a node the
    # renderer would interpret as markup.
    joined_text = " ".join(
        node.get("text", "")
        for page in pages
        for node in page.body["content"][0].get("content", [])
        if node.get("type") == "text"
    )
    assert "<script>" in joined_text or "script" in joined_text.lower()


async def test_markdown_heading_levels_are_clamped_to_schema_range(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="author4@example.com")
    markdown = "# Title\n\n## Section\n\n##### Too Deep\n\nBody text.\n"
    response = await _import_document(client, markdown.encode("utf-8"), filename="deep.md")
    assert response.status_code == 201, response.text
    body = response.json()
    assert "heading_level_adjusted" in _report_codes(body)

    pages = await _module_pages(db_session, body["module"]["id"])
    heading = next(node for node in pages[0].body["content"] if node["type"] == "heading")
    assert heading["attrs"]["level"] == 4


async def test_markdown_link_with_unsupported_scheme_is_dropped_to_plain_text(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="author5@example.com")
    markdown = "# Title\n\n## Section\n\nSee [our server](ftp://files.example.com/report).\n"
    response = await _import_document(client, markdown.encode("utf-8"), filename="link.md")
    assert response.status_code == 201, response.text
    body = response.json()
    assert "link_dropped" in _report_codes(body)

    pages = await _module_pages(db_session, body["module"]["id"])
    paragraph = pages[0].body["content"][0]
    assert not any(
        mark["type"] == "link"
        for node in paragraph.get("content", [])
        for mark in node.get("marks", [])
    )


async def test_markdown_import_is_rejected_when_not_utf8(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="author6@example.com")
    response = await _import_document(client, b"\xff\xfe\x00\x01", filename="bad.md")
    assert response.status_code == 422, response.text


async def test_document_import_audit_logged(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    author = await _author(client, db_session, email="author7@example.com")
    response = await _import_document(
        client, b"# Title\n\nBody text.\n", filename="basic.md"
    )
    assert response.status_code == 201, response.text
    entry = (
        await db_session.execute(
            select(AuditLog).where(AuditLog.action == "module_document_imported")
        )
    ).scalar_one()
    assert entry.actor_user_id == author.id
    assert entry.detail["source_format"] == "markdown"


async def test_document_import_requires_content_manager(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _login_with_role(
        client, db_session, email="learner@example.com", role_name="Learner"
    )
    response = await _import_document(client, b"# Title\n\nBody.\n", filename="basic.md")
    assert response.status_code == 403, response.text


# --- DOCX ----------------------------------------------------------------


async def test_docx_import_headings_lists_and_paragraphs(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="docxauthor1@example.com")
    response = await _import_document(
        client, _build_docx(), filename="policy.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["module"]["title"] == "Imported Title"

    pages = await _module_pages(db_session, body["module"]["id"])
    assert [page.title for page in pages] == ["Getting Started", "Next Section"]

    first_page_types = [node["type"] for node in pages[0].body["content"]]
    assert "bulletList" in first_page_types
    assert "orderedList" in first_page_types

    paragraph = pages[0].body["content"][0]
    assert any(
        mark["type"] == "bold"
        for node in paragraph["content"]
        for mark in node.get("marks", [])
    )


async def test_docx_embedded_image_produces_platform_hosted_asset(
    client: AsyncClient, db_session: AsyncSession, stored_objects: dict[str, Any]
) -> None:
    await _author(client, db_session, email="docxauthor2@example.com")
    response = await _import_document(
        client, _build_docx(with_image=True), filename="illustrated.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 201, response.text
    body = response.json()

    assets = (
        await db_session.execute(
            select(ModuleAsset).where(ModuleAsset.module_id == uuid.UUID(body["module"]["id"]))
        )
    ).scalars().all()
    assert len(assets) == 1
    assert assets[0].content_type == "image/png"
    assert assets[0].object_key in stored_objects

    pages = await _module_pages(db_session, body["module"]["id"])
    image_nodes = [
        node
        for page in pages
        for node in page.body["content"]
        if node["type"] == "image"
    ]
    assert len(image_nodes) == 1
    assert image_nodes[0]["attrs"]["src"] == (
        f"/api/modules/{body['module']['id']}/assets/{assets[0].id}"
    )


async def test_docx_table_is_dropped_and_reported(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="docxauthor3@example.com")
    response = await _import_document(
        client, _build_docx(with_table=True), filename="tabular.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert "table_dropped" in _report_codes(body)


async def test_docx_zip_bomb_is_rejected(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="docxauthor4@example.com")
    bomb = _build_zip(
        {
            "[Content_Types].xml": _WORD_CONTENT_TYPES,
            "word/document.xml": b"A" * (60 * 1024 * 1024),
        }
    )
    response = await _import_document(
        client, bomb, filename="bomb.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 413, response.text
    assert response.json()["detail"]["code"] == "suspicious_compression_ratio"

    modules = (await db_session.execute(select(Module))).scalars().all()
    assert modules == []


async def test_docx_renamed_spreadsheet_is_rejected(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="docxauthor5@example.com")
    fake = _build_zip(
        {
            "[Content_Types].xml": _SPREADSHEET_CONTENT_TYPES,
            "xl/workbook.xml": "<workbook/>",
        }
    )
    response = await _import_document(
        client, fake, filename="spreadsheet.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 422, response.text
    assert response.json()["detail"]["code"] == "unrecognised_format"


# --- PDF -----------------------------------------------------------------


async def test_pdf_import_extracts_text_and_warns_best_effort(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="pdfauthor1@example.com")
    response = await _import_document(
        client, _PDF_ONE_PAGE, filename="handout.pdf", content_type="application/pdf"
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert "pdf_best_effort" in _report_codes(body)
    assert "title_inferred_from_filename" in _report_codes(body)

    pages = await _module_pages(db_session, body["module"]["id"])
    assert len(pages) == 1
    text_content = pages[0].body["content"][0]["content"][0]["text"]
    assert "Hello world" in text_content


async def test_pdf_encrypted_is_rejected(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from pypdf import PdfWriter

    await _author(client, db_session, email="pdfauthor2@example.com")
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="secret")
    buffer = io.BytesIO()
    writer.write(buffer)

    response = await _import_document(
        client, buffer.getvalue(), filename="locked.pdf", content_type="application/pdf"
    )
    assert response.status_code == 422, response.text
    assert response.json()["detail"]["code"] == "encrypted_pdf"


# --- Format sniffing -------------------------------------------------------


async def test_unrecognised_format_is_rejected(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="fmtauthor1@example.com")
    response = await _import_document(
        client, b"\x7fELF\x02\x01\x01", filename="binary.bin"
    )
    assert response.status_code == 422, response.text
    assert response.json()["detail"]["code"] == "unrecognised_format"


async def test_extension_mismatch_is_rejected(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="fmtauthor2@example.com")
    response = await _import_document(
        client, _PDF_ONE_PAGE, filename="handout.docx", content_type="application/pdf"
    )
    assert response.status_code == 422, response.text
    assert response.json()["detail"]["code"] == "extension_mismatch"


async def test_docx_traversal_entry_name_is_rejected(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _author(client, db_session, email="docxauthor6@example.com")
    hostile = _build_zip(
        {
            "[Content_Types].xml": _WORD_CONTENT_TYPES,
            "../../etc/evil.xml": "<x/>",
        }
    )
    response = await _import_document(
        client, hostile, filename="traversal.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 422, response.text
    assert response.json()["detail"]["code"] == "unexpected_entry"

    modules = (await db_session.execute(select(Module))).scalars().all()
    assert modules == []


# --- Raw HTML sanitization (unreachable via the public import path with
# `html: False`, per the module docstring — exercised directly here, the same
# posture ticket 2's adversarial validator tests take for cases the editor UI
# itself would never produce). --------------------------------------------


def test_sanitize_raw_html_strips_script_content_entirely() -> None:
    from app.content.document_import import _sanitize_raw_html

    nodes = _sanitize_raw_html("<script>alert(1)</script>after")
    assert nodes == [{"type": "text", "text": "after"}]
    assert not any("script" in str(node).lower() for node in nodes)


def test_sanitize_raw_html_parses_supported_tags_into_marks() -> None:
    from app.content.document_import import _sanitize_raw_html

    nodes = _sanitize_raw_html('<b>bold</b> and <a href="https://example.com">link</a>')
    bold_node = next(node for node in nodes if node["text"] == "bold")
    assert bold_node["marks"] == [{"type": "bold"}]
    link_node = next(node for node in nodes if node["text"] == "link")
    assert link_node["marks"] == [{"type": "link", "attrs": {"href": "https://example.com"}}]


def test_sanitize_raw_html_drops_a_javascript_link_to_plain_text() -> None:
    from app.content.document_import import _sanitize_raw_html

    nodes = _sanitize_raw_html('<a href="javascript:alert(1)">bad</a>')
    assert nodes == [{"type": "text", "text": "bad"}]
